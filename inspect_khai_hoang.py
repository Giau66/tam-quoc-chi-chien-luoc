# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook('Meta Team Giới Thiệu Mùa PK.xlsx', data_only=True)
s = wb['Khai hoang']
starter_teams = []
cur_t = None
for r in range(20, s.max_row + 1):
    type_val = str(s.cell(r, 1).value or '').strip()
    gen_val = str(s.cell(r, 2).value or '').strip()
    cp1_val = str(s.cell(r, 3).value or '').strip()
    cp20_1 = str(s.cell(r, 5).value or '').strip()
    cp20_2 = str(s.cell(r, 6).value or '').strip()
    note = str(s.cell(r, 8).value or '').strip()
    
    # Check if starts a new team
    if any(t in type_val for t in ['Cung', 'Thương', 'Khiên', 'Kỵ']):
        if cur_t and len(cur_t['gens']) >= 2:
            starter_teams.append(cur_t)
        cur_t = {'type': type_val, 'gens': [], 'note': note}
    if cur_t and gen_val and gen_val not in ['Tướng', 'Võ Tướng']:
        cur_t['gens'].append({
            'name': gen_val,
            'cp_lv1': cp1_val,
            'cp_lv20': [c for c in [cp20_1, cp20_2] if c]
        })
        if note and not cur_t['note']:
            cur_t['note'] = note

if cur_t and len(cur_t['gens']) >= 2:
    starter_teams.append(cur_t)

print(f"Total starter teams parsed: {len(starter_teams)}")
for i, t in enumerate(starter_teams[:12]):
    names = [g['name'] for g in t['gens']]
    print(f"\nStarter {i+1} [{t['type']}]: {' - '.join(names)}")
    if t['note']:
        print(f"  Note: {t['note'][:80]}")
    for g in t['gens']:
        print(f"    {g['name']}: Đầu trận: {g['cp_lv1']} | Sau Lv20: {', '.join(g['cp_lv20'])}")
