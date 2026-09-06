import openpyxl, os, json

path = r'D:/TamQuocChiChienLuoc/Meta Team Giới Thiệu Mùa PK.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

out = {}
for sh in wb.sheetnames:
    ws = wb[sh]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 5:
            break
        rows.append([str(c) if c is not None else '' for c in row[:15]])
    out[sh] = rows

with open('tmpsheets.json', 'w', encoding='utf-8') as f:
    json.dump({'sheets': wb.sheetnames, 'preview': out}, f, ensure_ascii=False, indent=2)
print("Done")
