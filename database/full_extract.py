# -*- coding: utf-8 -*-
"""
Full extraction script: reads ALL 12 sheets from Excel and builds complete database.
"""
import sys
import openpyxl
import re
import os
import json
import unicodedata

sys.stdout.reconfigure(encoding='utf-8')
sys.stderr.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, "Meta Team Gioi Thieu Mua PK.xlsx")
DB_DIR = os.path.join(BASE_DIR, "database")

# Try both filenames (with and without diacritics)
if not os.path.exists(EXCEL_PATH):
    for fname in os.listdir(BASE_DIR):
        if fname.startswith("Meta Team") and fname.endswith(".xlsx"):
            EXCEL_PATH = os.path.join(BASE_DIR, fname)
            break

print(f"Excel path: {EXCEL_PATH}")
print(f"Exists: {os.path.exists(EXCEL_PATH)}")

def slugify(text):
    text = unicodedata.normalize('NFD', str(text))
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    text = text.lower().replace('\u0111', 'd').replace('\u0110', 'd')
    text = re.sub(r'[^a-z0-9]+', '_', text).strip('_')
    return text

GEN_NORM = {
    'gia hu': 'Gia Hu', 'tha tho': 'Tao Thao', 'lam thong': 'Lang Thong',
    'gi\u1ea3 h\u01b0': 'Gi\u1ea3 H\u1ee7', 'th\xe0o th\xe1o': 'T\xe0o Th\xe1o',
    'l\u0103m th\u1ed1ng': 'L\u0103ng Th\u1ed1ng',
    'sp qu\xe3n v\u0169': 'SP Quan V\u0169',
    'sp\u0111\u1ed3ng tr\xe1c': 'SP \u0110\u1ed3ng Tr\xe1c',
    'l\u1ee5c t\u1ed1\u1e47': 'L\u1ee5c T\u1ed1n',
    'b\xe1i s\u01b0 h\u1ee9a du': 'H\u1ee9a Du',
    'b\xe1i s\u01b0 l\u1ed7 t\xfac': 'L\u1ed7 T\xfac',
    'b\xe1i s\u01b0 ph\xe1p ch\xednh': 'Ph\xe1p Ch\xednh',
    'b\xe1i s\u01b0 sp tu\xe2n \xfac': 'SP Tu\xe2n \xdac',
    'l\u1eef b\u1ed1 40th': 'L\u1eef B\u1ed1',
    'co tinh th\xe1i': 'Tr\u01b0\u01a1ng Tinh Th\xe1i',
    'co \u0111\u1ea1i ki\u1ec1u': '\u0110\u1ea1i Ki\u1ec1u',
    'co ti\u1ec3u ki\u1ec1u': 'Ti\u1ec3u Ki\u1ec1u',
    'tinh th\xe1i': 'Tr\u01b0\u01a1ng Tinh Th\xe1i',
    't\u01b0 m\xe3 \xfd': 'T\u01b0 M\xe3 \xdd',
    'm\xe3n s\u1ee7ng': 'M\xe3n S\u1ee7ng',
    'tri\u1ec7u v\xe2n': 'Tri\u1ec7u V\xe2n',
    'kh\u01b0\u01a1ng duy': 'Kh\u01b0\u01a1ng Duy',
    'b\xe0ng th\u1ed1ng': 'B\xe0ng Th\u1ed1ng',
    'gia c\xe1t l\u01b0\u1ee3ng': 'Gia C\xe1t L\u01b0\u1ee3ng',
    't\xe0o th\xe1o': 'T\xe0o Th\xe1o',
    'chu th\xe1i': 'Chu Th\xe1i',
    'l\u0103ng th\u1ed1ng': 'L\u0103ng Th\u1ed1ng',
    'l\u01b0u b\u1ecb': 'L\u01b0u B\u1ecb',
    'tr\u01b0\u01a1ng phi': 'Tr\u01b0\u01a1ng Phi',
    'quan v\u0169': 'Quan V\u0169',
    'ng\u1ee5y di\xean': 'Ng\u1ee5y Di\xean',
    'quan ng\xe2n b\xecnh': 'Quan Ng\xe2n B\xecnh',
    'ph\xe1p ch\xednh': 'Ph\xe1p Ch\xednh',
    'sp m\xe3 si\xeau': 'SP M\xe3 Si\xeau',
    'sp ho\xe0ng ph\u1ed5 tung': 'SP Ho\xe0ng Ph\u1ed5 Tung',
    'sp tu\xe2n \xfac': 'SP Tu\xe2n \xdac',
    'sp qu\xe1ch gia': 'SP Qu\xe1ch Gia',
    'sp quan v\u0169': 'SP Quan V\u0169',
    'sp chu du': 'SP Chu Du',
    'sp l\u1eef m\xf4ng': 'SP L\u1eef M\xf4ng',
    'sp vi\xean thi\u1ec7u': 'SP Vi\xean Thi\u1ec7u',
    'sp chu tu\u1ea5n': 'SP Chu Tu\u1ea5n',
    'sp \u0111\u1ed3ng tr\xe1c': 'SP \u0110\u1ed3ng Tr\xe1c',
    'sp \u0111i\xeau thuy\u1ec1n': 'SP \u0110i\xeau Thuy\u1ec1n',
    'sp l\xf4 tr\u1ef1c': 'SP L\xf4 Tr\u1ef1c',
    'sp t\xf4n ki\xean': 'SP T\xf4n Ki\xean',
    'sp ho\xe0ng nguy\u1ec7t anh': 'SP Ho\xe0ng Nguy\u1ec7t Anh',
    'm\xe3 qu\xe2n': 'M\xe3 Qu\xe2n',
    'sp h\u1ee9a ch\u1eed': 'SP H\u1ee9a Ch\u1eed',
    'sp t\xe0o ch\xe2n': 'SP T\xe0o Ch\xe2n',
    'd\u01b0\u01a1ng d\u1ecb': 'D\u01b0\u01a1ng D\u1ecb',
    'sp tr\u01b0\u01a1ng b\u1ea3o': 'SP Tr\u01b0\u01a1ng B\u1ea3o',
    'sp tr\u01b0\u01a1ng l\u01b0\u01a1ng': 'SP Tr\u01b0\u01a1ng L\u01b0\u01a1ng',
    'sp b\xe0ng \u0111\u1ee9c': 'SP B\xe0ng \u0110\u1ee9c',
    'sp l\u01b0u di\u1ec7p': 'SP L\u01b0u Di\u1ec7p',
    'gia c\xe1t kh\xe1c': 'Gia C\xe1t Kh\xe1c',
    'm\xe3 \u0111\u1ea1i': 'M\xe3 \u0110\u1ea1i',
    'sp ph\xe1p ch\xednh': 'SP Ph\xe1p Ch\xednh'
}

FACTION_MAP = {
    'L\u01b0u B\u1ecb': 'Th\u1ee5c', 'Quan V\u0169': 'Th\u1ee5c', 'Tr\u01b0\u01a1ng Phi': 'Th\u1ee5c', 'Tri\u1ec7u V\xe2n': 'Th\u1ee5c',
    'Gia C\xe1t L\u01b0\u1ee3ng': 'Th\u1ee5c', 'B\xe0ng Th\u1ed1ng': 'Th\u1ee5c', 'Kh\u01b0\u01a1ng Duy': 'Th\u1ee5c', 'Ng\u1ee5y Di\xean': 'Th\u1ee5c',
    'Ho\xe0ng Trung': 'Th\u1ee5c', 'M\xe3 Si\xeau': 'Th\u1ee5c', 'Ph\xe1p Ch\xednh': 'Th\u1ee5c', 'Quan Ng\xe2n B\xecnh': 'Th\u1ee5c',
    'Tr\u01b0\u01a1ng Tinh Th\xe1i': 'Th\u1ee5c', 'Ho\xe0ng Nguy\u1ec7t Anh': 'Th\u1ee5c', 'Nghi\xeam Nhan': 'Th\u1ee5c',
    'M\xe3 V\xe2n L\u1ed9c': 'Th\u1ee5c', 'V\u01b0\u01a1ng B\xecnh': 'Th\u1ee5c', 'T\u1eeb Th\u1ee9': 'Th\u1ee5c', 'SP Quan V\u0169': 'Th\u1ee5c',
    'SP Gia C\xe1t L\u01b0\u1ee3ng': 'Th\u1ee5c', 'SP Ho\xe0ng Nguy\u1ec7t Anh': 'Th\u1ee5c', 'SP M\xe3 Si\xeau': 'Th\u1ee5c',
    'M\xe3 \u0110\u1ea1i': 'Th\u1ee5c', 'SP Ph\xe1p Ch\xednh': 'Th\u1ee5c', 'Li\xeau H\xf3a': 'Th\u1ee5c',
    'T\xe0o Th\xe1o': 'Ng\u1ee5y', 'T\u01b0 M\xe3 \xdd': 'Ng\u1ee5y', 'Qu\xe1ch Gia': 'Ng\u1ee5y', 'Gi\u1ea3 H\u1ee7': 'Ng\u1ee5y',
    'Tr\xecnh D\u1ee5c': 'Ng\u1ee5y', 'H\u1ea1 H\u1ea7u Uy\xean': 'Ng\u1ee5y', 'H\u1ea1 H\u1ea7u \u0110\xf4n': 'Ng\u1ee5y', 'Tr\u01b0\u01a1ng Li\xeau': 'Ng\u1ee5y',
    'H\xe1c Chi\xeau': 'Ng\u1ee5y', 'M\xe3n S\u1ee7ng': 'Ng\u1ee5y', 'Tu\xe2n \xdac': 'Ng\u1ee5y', 'Tu\xe2n Du': 'Ng\u1ee5y',
    'V\u01b0\u01a1ng Nguy\xean C\u01a1': 'Ng\u1ee5y', 'Chung H\u1ed9i': 'Ng\u1ee5y', '\u0110\u1eb7ng Ng\u1ea3i': 'Ng\u1ee5y', 'H\u1ee9a Ch\u1eed': 'Ng\u1ee5y',
    '\u0110i\u1ec3n Vi': 'Ng\u1ee5y', 'T\xe0o Nh\xe2n': 'Ng\u1ee5y', 'T\xe0o Phi': 'Ng\u1ee5y', 'T\xe0o Thu\u1ea7n': 'Ng\u1ee5y',
    'B\xe0ng \u0110\u1ee9c': 'Ng\u1ee5y', 'Ch\xe2n C\u01a1': 'Ng\u1ee5y', 'SP Tu\xe2n \xdac': 'Ng\u1ee5y', 'SP Qu\xe1ch Gia': 'Ng\u1ee5y',
    'SP H\u1ee9a Ch\u1eed': 'Ng\u1ee5y', 'SP T\xe0o Ch\xe2n': 'Ng\u1ee5y', 'SP B\xe0ng \u0110\u1ee9c': 'Ng\u1ee5y', 'SP L\u01b0u Di\u1ec7p': 'Ng\u1ee5y',
    'T\xf4n Quy\u1ec1n': 'Ng\xf4', 'T\xf4n S\xe1ch': 'Ng\xf4', 'T\xf4n Ki\xean': 'Ng\xf4', 'T\xf4n Th\u01b0\u1ee3ng H\u01b0\u01a1ng': 'Ng\xf4',
    'Chu Du': 'Ng\xf4', 'L\u1ee5c T\u1ed1n': 'Ng\xf4', 'L\u1ed7 T\xfac': 'Ng\xf4', 'L\u1eef M\xf4ng': 'Ng\xf4',
    'Th\xe1i S\u1eed T\u1eeb': 'Ng\xf4', 'Cam Ninh': 'Ng\xf4', 'Tr\xecnh Ph\u1ed5': 'Ng\xf4', 'Ho\xe0ng C\xe1i': 'Ng\xf4',
    'Chu Th\xe1i': 'Ng\xf4', 'L\u0103ng Th\u1ed1ng': 'Ng\xf4', '\u0110\u1ea1i Ki\u1ec1u': 'Ng\xf4', 'Ti\u1ec3u Ki\u1ec1u': 'Ng\xf4',
    'L\u1ee5c Kh\xe1ng': 'Ng\xf4', '\u0110inh Ph\u1ee5ng': 'Ng\xf4', 'H\xe0n \u0110\u01b0\u01a1ng': 'Ng\xf4', 'T\u01b0\u1edfng Kh\xe2m': 'Ng\xf4',
    'SP Chu Du': 'Ng\xf4', 'SP L\u1eef M\xf4ng': 'Ng\xf4', 'SP T\xf4n Ki\xean': 'Ng\xf4', 'Gia C\xe1t Kh\xe1c': 'Ng\xf4',
    'L\u1eef B\u1ed1': 'Qu\u1ea7n', '\u0110i\xeau Thuy\u1ec1n': 'Qu\u1ea7n', '\u0110\u1ed3ng Tr\xe1c': 'Qu\u1ea7n', 'Hoa \u0110\xe0': 'Qu\u1ea7n',
    'T\u1ea3 T\u1eeb': 'Qu\u1ea7n', 'Vu C\xe1t': 'Qu\u1ea7n', 'Tr\u01b0\u01a1ng Gi\xe1c': 'Qu\u1ea7n', 'Vi\xean Thi\u1ec7u': 'Qu\u1ea7n',
    'Chu Tu\u1ea5n': 'Qu\u1ea7n', 'Th\u01b0 Th\u1ee5': 'Qu\u1ea7n', 'Hoa H\xf9ng': 'Qu\u1ea7n', 'Nhan L\u01b0\u01a1ng': 'Qu\u1ea7n',
    'V\u0103n X\xfa': 'Qu\u1ea7n', 'Ch\xfac Dung': 'Qu\u1ea7n', 'M\u1ea1nh Ho\u1ea1ch': 'Qu\u1ea7n', 'Tr\u1ea7n Cung': 'Qu\u1ea7n',
    '\u0110i\u1ec1n Phong': 'Qu\u1ea7n', 'L\xfd Nho': 'Qu\u1ea7n', 'H\u1ee9a Du': 'Qu\u1ea7n', 'D\u01b0\u01a1ng D\u1ecb': 'Qu\u1ea7n',
    'M\xe3 Qu\xe2n': 'Qu\u1ea7n', 'SP Vi\xean Thi\u1ec7u': 'Qu\u1ea7n', 'SP Chu Tu\u1ea5n': 'Qu\u1ea7n', 'SP \u0110\u1ed3ng Tr\xe1c': 'Qu\u1ea7n',
    'SP \u0110i\xeau Thuy\u1ec1n': 'Qu\u1ea7n', 'SP L\xf4 Tr\u1ef1c': 'Qu\u1ea7n', 'SP Ho\xe0ng Ph\u1ed5 Tung': 'Qu\u1ea7n',
    'SP Tr\u01b0\u01a1ng B\u1ea3o': 'Qu\u1ea7n', 'SP Tr\u01b0\u01a1ng L\u01b0\u01a1ng': 'Qu\u1ea7n', 'Kh\xfac Ngh\u0129a': 'Qu\u1ea7n',
    'Cao L\xe3m': 'Qu\u1ea7n', 'Chu Thi\u1ec7u': 'Qu\u1ea7n', 'Vi\xean Thu\u1eadt': 'Qu\u1ea7n'
}

def clean_gen_name(name):
    if not name:
        return ""
    name = str(name).strip()
    name = re.sub(r'[\r\n]+', ' ', name)
    name = re.sub(r'\([^\)]*\)', '', name)
    name = re.sub(r'\b(40th|CO|Ch\u1ee7 T\u01b0\u1edbng|Ph\xf3 T\u01b0\u1edbng)\b', '', name, flags=re.I)
    name = name.strip()
    name = re.sub(r'\s+', ' ', name)
    n_lower = name.lower()
    return GEN_NORM.get(n_lower, name)

def clean_tac_name(name):
    if not name:
        return ""
    name = str(name).strip()
    name = re.sub(r'[\r\n]+', ' ', name)
    name = re.sub(r'^(Tinh|Chuy\u1ec3n|Th\xe1c)[·\.\s]', '', name)
    name = re.sub(r'\([^\)]*\)', '', name)
    name = name.strip()
    name = re.sub(r'\s+', ' ', name)
    return name

def infer_faction(name):
    clean = clean_gen_name(name)
    if clean in FACTION_MAP:
        return FACTION_MAP[clean]
    for k, v in FACTION_MAP.items():
        if clean == k or clean in k:
            return v
    return "Qu\u1ea7n"

def infer_tactic_type(name):
    n = name.lower()
    if 'tr\u1eadn' in n:
        return 'Tr\u1eadn ph\xe1p'
    if any(k in n for k in ['k\u1ef5', 'khi\xean', 'cung', 'th\u01b0\u01a1ng', 'binh', 'doanh', 'v\u1ec7', 's\u0129']):
        return 'Binh ch\u1ee7ng'
    if any(k in n for k in ['th\xe1i b\xecnh', 's\u0129 bi\u1ec7t', 'd\u1ee5ng v\xf5', 'tuy\u1ec7t \u0111\u1ecba', 'h\u1ed5 c\u1ee9', 'l\xf5a y', 'binh v\xf4', 'v\u0103n thao']):
        return 'B\u1ecb \u0111\u1ed9ng'
    if any(k in n for k in ['b\xe1t m\xf4n', 'qu\xe2n d\xe2n', 'th\u1ea3o thuy\u1ec1n', 't\u1ea1m th\u1eddi', 'th\u1ecbnh kh\xed', 'ti\u1ec1m long', 'c\u01a1 h\xecnh', 'gi\u1ea3i phi\u1ec1n']):
        return 'Ch\u1ec9 huy'
    if any(k in n for k in ['nh\u1ea5t k\u1ef5', 'qu\u1ef7 th\u1ea7n', 'b\xe1ch k\u1ef5', 'b\u1ea1o l\u1ec7', '\u0111\u01b0\u01a1ng phong', 'mau gi\xe0nh']):
        return '\u0110\u1ed9t k\xedch'
    return 'Ch\u1ee7 \u0111\u1ed9ng'

def main():
    print(f"Loading Excel from: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    print(f"Sheets: {wb.sheetnames}")

    generals_data = {}
    tactics_data = {}
    meta_teams_list = []
    coexisting_portfolios = []
    team_do_uy_data = []
    mua4_teams = []
    mua5_teams = []
    cot_truyen_data = []
    khai_hoang_data = []
    cai_tao_data = []

    def reg_gen(name, season="PK"):
        if not name or name in generals_data:
            return
        f = infer_faction(name)
        gid = f"{slugify(f)}_{slugify(name)}"
        generals_data[name] = {
            "id": gid, "name": name, "faction": f,
            "cost": 7 if 'SP' in name else 6,
            "troop": {"Ky": "A", "Khien": "A", "Cung": "A", "Thuong": "A", "Khi": "B"},
            "inherent_skill": f"Ky nang rieng cua {name}",
            "role": ["Tuong chu luc / Ho tro trong Meta"],
            "season": season
        }

    def reg_tac(name, desc="", season="PK"):
        if not name or name in tactics_data:
            return
        tactics_data[name] = {
            "id": slugify(name), "name": name,
            "type": infer_tactic_type(name),
            "quality": "S",
            "description": desc or f"Chien phap manh trong meta {season}",
            "season": season
        }

    # ==========================================================
    # 1. Tuong&CP
    # ==========================================================
    sheet_name = None
    for sn in wb.sheetnames:
        if 'ng&CP' in sn or 'ong&CP' in sn or ('ng' in sn and 'CP' in sn):
            sheet_name = sn
            break
    if sheet_name:
        s_cp = wb[sheet_name]
        mode = None
        season = "PK"
        for r in range(1, s_cp.max_row + 1):
            c1 = str(s_cp.cell(r, 1).value or '').strip()
            c2 = str(s_cp.cell(r, 2).value or '').strip()
            c3 = str(s_cp.cell(r, 3).value or '').strip()
            c4 = str(s_cp.cell(r, 4).value or '').strip()
            c5 = str(s_cp.cell(r, 5).value or '').strip()
            if 'PK' in c1 or 'mua' in c1.lower() or 'M\u00f9a' in c1:
                season = c1
            if 'Tuong Moi' in c1 or 'T\u01b0\u1edbng M\u1edbi' in c1 or 'Tuong Moi' in c2 or 'T\u01b0\u1edbng M\u1edbi' in c2:
                mode = 'GEN'
                continue
            elif 'Chien Phap Su Kien' in c1 or 'Chi\u1ebfn Ph\u00e1p S\u1ef1 Ki\u1ec7n' in c1 or \
                 'Chi\u1ebfn Ph\u00e1p S\u1ef1 Ki\u1ec7n' in c2 or 'Chi\u1ebfn Ph\u00e1p K\u1ebf Th\u1eeba' in c2:
                mode = 'TAC'
                continue
            if c1 in ['STT', ''] and c2 in ['T\xean T\u01b0\u1edbng', 'T\xean Chi\u1ebfn Ph\u00e1p']:
                continue
            if mode == 'GEN' and c2:
                g_name = clean_gen_name(c2)
                if g_name:
                    f = infer_faction(g_name)
                    gid = f"{slugify(f)}_{slugify(g_name)}"
                    generals_data[g_name] = {
                        "id": gid, "name": g_name, "faction": f,
                        "cost": 7 if 'SP' in g_name else 6,
                        "troop": {"Ky": "S", "Khien": "A", "Cung": "S", "Thuong": "S", "Khi": "A"},
                        "inherent_skill": c3 if c3 else f"Thien phu cua {g_name}",
                        "inherit_tactic": clean_tac_name(c4) if c4 != '-' else '',
                        "role": [c5] if c5 and c5 != '-' else ["Chu luc / Ho tro chien thuat"],
                        "season": season
                    }
                    if c4 and c4 != '-':
                        t_name = clean_tac_name(c4)
                        if t_name:
                            reg_tac(t_name, f"Chien phap ke thua tu {g_name}", season)
            elif mode == 'TAC' and c2:
                t_name = clean_tac_name(c2)
                if t_name:
                    desc = c3
                    if c5:
                        desc += f" (Can: {c5})"
                    tactics_data[t_name] = {
                        "id": slugify(t_name), "name": t_name,
                        "type": infer_tactic_type(t_name),
                        "quality": "S", "description": desc,
                        "suitable_generals": c4 if c4 else "",
                        "season": season
                    }
        print(f"[Tuong&CP] Generals: {len(generals_data)}, Tactics: {len(tactics_data)}")

    # ==========================================================
    # 2. Cot truyen Moi
    # ==========================================================
    for sn in wb.sheetnames:
        if 'truyen' in sn.lower() or 'truy\u1ec7n' in sn.lower():
            s_ct = wb[sn]
            for r in range(2, s_ct.max_row + 1):
                c1 = str(s_ct.cell(r, 1).value or '').strip()
                c2 = str(s_ct.cell(r, 2).value or '').strip()
                c3 = str(s_ct.cell(r, 3).value or '').strip()
                c4 = str(s_ct.cell(r, 4).value or '').strip()
                if c2:
                    cot_truyen_data.append({"stt": c1, "ten_tuong": c2, "phe": c3, "mo_ta": c4})
            print(f"[Cot truyen] Entries: {len(cot_truyen_data)}")
            break

    # ==========================================================
    # 3. Khai hoang
    # ==========================================================
    for sn in wb.sheetnames:
        if 'khai' in sn.lower() or 'hoang' in sn.lower():
            s_kh = wb[sn]
            rows_data = []
            for r in range(1, s_kh.max_row + 1):
                rv = [str(s_kh.cell(r, c).value or '').strip() for c in range(1, 20)]
                rows_data.append(rv)
            khai_hoang_data = {"sheet": sn, "rows": rows_data[:200]}
            print(f"[Khai hoang] Rows collected: {len(rows_data)}")
            break

    # ==========================================================
    # 4. Team gioi thieu (moi) - PRIMARY SHEET
    # ==========================================================
    for sn in wb.sheetnames:
        if 'gi\u1edbi thi\u1ec7u' in sn.lower() or 'gioi thieu' in sn.lower():
            s_new = wb[sn]
            cur_t = None
            for r in range(5, s_new.max_row + 1):
                type_val = str(s_new.cell(r, 1).value or '').strip()
                gen_val  = str(s_new.cell(r, 2).value or '').strip()
                cp1_val  = str(s_new.cell(r, 3).value or '').strip()
                cp2_val  = str(s_new.cell(r, 4).value or '').strip()
                bt1 = str(s_new.cell(r, 5).value or '').strip()
                bt2 = str(s_new.cell(r, 6).value or '').strip()
                bt3 = str(s_new.cell(r, 7).value or '').strip()
                note = str(s_new.cell(r, 8).value or '').strip()
                if type_val:
                    if cur_t and len(cur_t['generals']) >= 2:
                        meta_teams_list.append(cur_t)
                    troop = "Ky"
                    for t_name in ["Thuong", "Khien", "Cung", "Ky", "Khi",
                                   "Th\u01b0\u01a1ng", "Khi\xean", "K\u1ef5", "Kh\xed"]:
                        if t_name in type_val:
                            troop = t_name
                            break
                    tier = "T0"
                    if "T0.5" in type_val or "3" in type_val:
                        tier = "T0.5"
                    elif "T1" in type_val:
                        tier = "T1"
                    elif "T2" in type_val:
                        tier = "T2"
                    cur_t = {
                        "source": sn, "troop": troop, "tier": tier,
                        "note": note, "season": "PK", "generals": []
                    }
                if cur_t and gen_val:
                    cur_t['generals'].append({
                        "raw_name": gen_val,
                        "cp1_raw": cp1_val, "cp2_raw": cp2_val,
                        "binh_thu": [b for b in [bt1, bt2, bt3] if b]
                    })
                    if note and not cur_t['note']:
                        cur_t['note'] = note
            if cur_t and len(cur_t['generals']) >= 2:
                meta_teams_list.append(cur_t)
            print(f"[{sn}] Raw teams: {len(meta_teams_list)}")
            break

    # ==========================================================
    # 5. Cai Tao Binh Chung
    # ==========================================================
    for sn in wb.sheetnames:
        if 'Cai Tao' in sn or 'C\u1ea3i T\u1ea1o' in sn:
            s_ct2 = wb[sn]
            rows_data = []
            for r in range(1, s_ct2.max_row + 1):
                rv = [str(s_ct2.cell(r, c).value or '').strip() for c in range(1, 15)]
                if any(rv):
                    rows_data.append(rv)
            cai_tao_data = {"sheet": sn, "rows": rows_data[:300]}
            print(f"[Cai Tao BC] Rows: {len(rows_data)}")
            break

    # ==========================================================
    # 6. Team Do Uy
    # ==========================================================
    for sn in wb.sheetnames:
        if '\u0110\xf4' in sn and 'y' in sn and 'Team' in sn and 'gi\u1edbi' not in sn:
            s_du = wb[sn]
            cur_t = None
            TROOPS = ['\u0110\u1ea1i Kh\xedch', 'Th\u01b0\u01a1ng', 'Khi\xean', 'Cung', 'K\u1ef5', 'Kh\xed']
            for r in range(6, s_du.max_row + 1):
                type_val = str(s_du.cell(r, 1).value or '').strip()
                gen_val  = str(s_du.cell(r, 2).value or '').strip()
                cp1_val  = str(s_du.cell(r, 3).value or '').strip()
                cp2_val  = str(s_du.cell(r, 4).value or '').strip()
                bt1 = str(s_du.cell(r, 5).value or '').strip()
                bt2 = str(s_du.cell(r, 6).value or '').strip()
                bt3 = str(s_du.cell(r, 7).value or '').strip()
                note = str(s_du.cell(r, 8).value or '').strip() if s_du.max_column >= 8 else ''
                is_troop_row = type_val and any(t in type_val for t in TROOPS)
                if is_troop_row:
                    if cur_t and len(cur_t['generals']) >= 2:
                        team_do_uy_data.append(cur_t)
                    troop = type_val.strip()
                    cur_t = {
                        "source": sn, "troop": troop, "tier": "T1",
                        "note": note, "season": "PK", "generals": []
                    }
                if cur_t and gen_val:
                    g_name = clean_gen_name(gen_val)
                    if g_name and g_name not in ['Type', 'T\u01b0\u1edbng', 'V\xf5 T\u01b0\u1edbng']:
                        cur_t['generals'].append({
                            "raw_name": gen_val, "cp1_raw": cp1_val, "cp2_raw": cp2_val,
                            "binh_thu": [b for b in [bt1, bt2, bt3] if b]
                        })
                        if note and not cur_t['note']:
                            cur_t['note'] = note
            if cur_t and len(cur_t['generals']) >= 2:
                team_do_uy_data.append(cur_t)
            print(f"[Team Do Uy] Teams: {len(team_do_uy_data)}")
            break


    # ==========================================================
    # 7. Team Cung Ton
    # ==========================================================
    for sn in wb.sheetnames:
        if 'C\xf9ng T\u1ed3n' in sn or 'Cung Ton' in sn or ('Cung' in sn and 'Ton' in sn):
            s_tc = wb[sn]
            cur_set = None
            cur_sub_team = None
            for r in range(4, s_tc.max_row + 1):
                c1 = str(s_tc.cell(r, 1).value or '').strip()
                c2 = str(s_tc.cell(r, 2).value or '').strip()
                c3 = str(s_tc.cell(r, 3).value or '').strip()
                c4 = str(s_tc.cell(r, 4).value or '').strip()
                c5 = str(s_tc.cell(r, 5).value or '').strip()
                if 'C\xf9ng t\u1ed3n' in c1 or 'cung ton' in c1.lower():
                    if cur_set:
                        coexisting_portfolios.append(cur_set)
                    cur_set = {"set_name": c1, "description": c2, "teams": []}
                    cur_sub_team = None
                    continue
                if c2 and any(t in c2 for t in ['Th\u01b0\u01a1ng', 'Khi\xean', 'Cung', 'K\u1ef5', 'Kh\xed', '19C', '20C', '18C', '16C']):
                    cur_sub_team = {"team_name": c2, "generals": []}
                    if cur_set:
                        cur_set["teams"].append(cur_sub_team)
                if cur_sub_team and c3:
                    g_c = clean_gen_name(c3)
                    if g_c and g_c not in ['V\xf5 T\u01b0\u1edbng', 'T\u01b0\u1edbng']:
                        cur_sub_team["generals"].append({
                            "name": g_c,
                            "tactic1": clean_tac_name(c4),
                            "tactic2": clean_tac_name(c5)
                        })
                        reg_gen(g_c)
            if cur_set:
                coexisting_portfolios.append(cur_set)
            print(f"[Team Cung Ton] Sets: {len(coexisting_portfolios)}")
            break

    # ==========================================================
    # 8. Meta moi phe (cu)
    # ==========================================================
    for sn in wb.sheetnames:
        if 'moi phe' in sn.lower() or 'm\u1ed7i phe' in sn.lower():
            s_phe = wb[sn]
            cur_t = None
            for r in range(4, s_phe.max_row + 1):
                type_val = str(s_phe.cell(r, 1).value or '').strip()
                gen_val  = str(s_phe.cell(r, 2).value or '').strip()
                cp1_val  = str(s_phe.cell(r, 3).value or '').strip()
                cp2_val  = str(s_phe.cell(r, 4).value or '').strip()
                note     = str(s_phe.cell(r, 8).value or '').strip() if s_phe.max_column >= 8 else ''
                if type_val:
                    if cur_t and len(cur_t['generals']) >= 2:
                        meta_teams_list.append(cur_t)
                    troop = "Ky"
                    for t in ["Th\u01b0\u01a1ng", "Khi\xean", "Cung", "K\u1ef5", "Kh\xed"]:
                        if t in type_val:
                            troop = t
                            break
                    cur_t = {
                        "source": sn, "troop": troop, "tier": "T1",
                        "note": note, "season": "PK", "generals": []
                    }
                if cur_t and gen_val:
                    cur_t['generals'].append({
                        "raw_name": gen_val, "cp1_raw": cp1_val, "cp2_raw": cp2_val, "binh_thu": []
                    })
                    if note and not cur_t['note']:
                        cur_t['note'] = note
            if cur_t and len(cur_t['generals']) >= 2:
                meta_teams_list.append(cur_t)
            break

    # ==========================================================
    # 9. Meta Team (official)
    # ==========================================================
    for sn in wb.sheetnames:
        if sn.strip() == 'Meta Team':
            s_meta = wb[sn]
            cur_t = None
            for r in range(3, s_meta.max_row + 1):
                stt      = str(s_meta.cell(r, 1).value or '').strip()
                type_val = str(s_meta.cell(r, 2).value or '').strip()
                power    = str(s_meta.cell(r, 3).value or '').strip()
                gen_val  = str(s_meta.cell(r, 4).value or '').strip()
                cp1_val  = str(s_meta.cell(r, 5).value or '').strip()
                cp2_val  = str(s_meta.cell(r, 6).value or '').strip()
                note     = str(s_meta.cell(r, 10).value or '').strip() if s_meta.max_column >= 10 else ''
                if stt or type_val:
                    if cur_t and len(cur_t['generals']) >= 2:
                        meta_teams_list.append(cur_t)
                    troop = "Ky"
                    for t in ["Th\u01b0\u01a1ng", "Khi\xean", "Cung", "K\u1ef5", "Kh\xed"]:
                        if t in type_val:
                            troop = t
                            break
                    cur_t = {
                        "source": sn, "troop": troop,
                        "tier": power if power else "T1",
                        "note": note, "season": "PK", "generals": []
                    }
                if cur_t and gen_val:
                    cur_t['generals'].append({
                        "raw_name": gen_val, "cp1_raw": cp1_val, "cp2_raw": cp2_val, "binh_thu": []
                    })
                    if note and not cur_t['note']:
                        cur_t['note'] = note
            if cur_t and len(cur_t['generals']) >= 2:
                meta_teams_list.append(cur_t)
            break

    # ==========================================================
    # 10. nhieu Team cung ton tai
    # ==========================================================
    for sn in wb.sheetnames:
        if 'nhi\u1ec1u' in sn.lower() or 'nhieu' in sn.lower():
            s_nhieu = wb[sn]
            cur_set = None
            cur_sub_team = None
            for r in range(3, s_nhieu.max_row + 1):
                c1 = str(s_nhieu.cell(r, 1).value or '').strip()
                c2 = str(s_nhieu.cell(r, 2).value or '').strip()
                c3 = str(s_nhieu.cell(r, 3).value or '').strip()
                c4 = str(s_nhieu.cell(r, 4).value or '').strip()
                c5 = str(s_nhieu.cell(r, 5).value or '').strip()
                if c1 and ('C\xf9ng t\u1ed3n' in c1 or 'cung ton' in c1.lower() or re.match(r'\d+\.', c1)):
                    if cur_set:
                        coexisting_portfolios.append(cur_set)
                    cur_set = {"set_name": c1, "description": c2, "teams": [], "source": sn}
                    cur_sub_team = None
                    continue
                if c2 and any(t in c2 for t in ['Th\u01b0\u01a1ng', 'Khi\xean', 'Cung', 'K\u1ef5', 'Kh\xed']):
                    cur_sub_team = {"team_name": c2, "generals": []}
                    if cur_set:
                        cur_set["teams"].append(cur_sub_team)
                if cur_sub_team and c3:
                    g_c = clean_gen_name(c3)
                    if g_c and g_c not in ['V\xf5 T\u01b0\u1edbng', 'T\u01b0\u1edbng']:
                        cur_sub_team["generals"].append({
                            "name": g_c, "tactic1": clean_tac_name(c4), "tactic2": clean_tac_name(c5)
                        })
                        reg_gen(g_c)
            if cur_set:
                coexisting_portfolios.append(cur_set)
            break

    # ==========================================================
    # 11. Meta team mua 4
    # ==========================================================
    for sn in wb.sheetnames:
        if 'm\xf9a 4' in sn.lower() or 'mua 4' in sn.lower():
            s_m4 = wb[sn]
            cur_t = None
            for r in range(3, s_m4.max_row + 1):
                type_val = str(s_m4.cell(r, 1).value or '').strip()
                gen_val  = str(s_m4.cell(r, 2).value or '').strip()
                cp1_val  = str(s_m4.cell(r, 3).value or '').strip()
                cp2_val  = str(s_m4.cell(r, 4).value or '').strip()
                if type_val and any(t in type_val for t in ['Th\u01b0\u01a1ng', 'Khi\xean', 'Cung', 'K\u1ef5', 'Kh\xed']):
                    if cur_t and len(cur_t['generals']) >= 2:
                        mua4_teams.append(cur_t)
                    troop = "Ky"
                    for t in ["Th\u01b0\u01a1ng", "Khi\xean", "Cung", "K\u1ef5", "Kh\xed"]:
                        if t in type_val:
                            troop = t
                            break
                    cur_t = {
                        "source": sn, "troop": troop, "tier": "T1",
                        "note": "", "season": "Mua 4", "generals": []
                    }
                if cur_t and gen_val:
                    cur_t['generals'].append({
                        "raw_name": gen_val, "cp1_raw": cp1_val, "cp2_raw": cp2_val, "binh_thu": []
                    })
            if cur_t and len(cur_t['generals']) >= 2:
                mua4_teams.append(cur_t)
            print(f"[{sn}] Teams: {len(mua4_teams)}")
            break

    # ==========================================================
    # 12. Meta team mua 5
    # ==========================================================
    for sn in wb.sheetnames:
        if 'm\xf9a 5' in sn.lower() or 'mua 5' in sn.lower():
            s_m5 = wb[sn]
            cur_t = None
            for r in range(3, s_m5.max_row + 1):
                type_val = str(s_m5.cell(r, 1).value or '').strip()
                gen_val  = str(s_m5.cell(r, 2).value or '').strip()
                cp1_val  = str(s_m5.cell(r, 3).value or '').strip()
                cp2_val  = str(s_m5.cell(r, 4).value or '').strip()
                if type_val and any(t in type_val for t in ['Th\u01b0\u01a1ng', 'Khi\xean', 'Cung', 'K\u1ef5', 'Kh\xed']):
                    if cur_t and len(cur_t['generals']) >= 2:
                        mua5_teams.append(cur_t)
                    troop = "Ky"
                    for t in ["Th\u01b0\u01a1ng", "Khi\xean", "Cung", "K\u1ef5", "Kh\xed"]:
                        if t in type_val:
                            troop = t
                            break
                    cur_t = {
                        "source": sn, "troop": troop, "tier": "T1",
                        "note": "", "season": "Mua 5", "generals": []
                    }
                if cur_t and gen_val:
                    cur_t['generals'].append({
                        "raw_name": gen_val, "cp1_raw": cp1_val, "cp2_raw": cp2_val, "binh_thu": []
                    })
            if cur_t and len(cur_t['generals']) >= 2:
                mua5_teams.append(cur_t)
            print(f"[{sn}] Teams: {len(mua5_teams)}")
            break

    # ==========================================================
    # PROCESS & DEDUPLICATE
    # ==========================================================
    all_raw = meta_teams_list + mua4_teams + mua5_teams
    processed_teams = []
    seen_keys = set()
    counter = 1
    positions = ["Chu tuong", "Pho tuong 1", "Pho tuong 2"]

    for raw_t in all_raw:
        gens = raw_t.get('generals', [])
        if len(gens) < 2:
            continue
        team_gens = []
        gen_names_key = []
        for idx, g_obj in enumerate(gens[:3]):
            raw_name = g_obj['raw_name']
            alts = [clean_gen_name(p) for p in raw_name.split('/') if clean_gen_name(p)]
            if not alts:
                continue
            mn = alts[0]
            alt_gens = alts[1:]
            gen_names_key.append(mn)
            reg_gen(mn, raw_t.get('season', 'PK'))
            bis = []
            sub = []
            for cp_raw in [g_obj.get('cp1_raw', ''), g_obj.get('cp2_raw', '')]:
                if not cp_raw:
                    continue
                parts = [clean_tac_name(p) for p in re.split(r'[/,\n]+', cp_raw) if clean_tac_name(p)]
                if parts:
                    bis.append(parts[0])
                    reg_tac(parts[0], "", raw_t.get('season', 'PK'))
                    for alt_t in parts[1:]:
                        sub.append(alt_t)
                        reg_tac(alt_t, "", raw_t.get('season', 'PK'))
            f = infer_faction(mn)
            gid = generals_data[mn]["id"]
            team_gens.append({
                "position": positions[idx] if idx < len(positions) else f"Pho tuong {idx}",
                "general_id": gid,
                "name": mn,
                "bis_tactics": bis,
                "sub_tactics": sub,
                "binh_thu": g_obj.get('binh_thu', []),
                "alt_generals": alt_gens
            })
        if len(team_gens) < 2:
            continue
        key = "_".join(sorted(gen_names_key)) + f"_{raw_t['troop']}"
        if key in seen_keys:
            continue
        seen_keys.add(key)
        factions = [infer_faction(g['name']) for g in team_gens]
        team_faction = factions[0] if len(set(factions)) == 1 else "Tam The"
        notes_str = raw_t.get('note', '')
        strengths = []
        weaknesses = []
        if "Cham:" in notes_str or "Ch\u1ea1m:" in notes_str:
            key_str = "Ch\u1ea1m:" if "Ch\u1ea1m:" in notes_str else "Cham:"
            match_part = notes_str.split(key_str)[1]
            avoid_key = "Tr\u00e1nh:" if "Tr\u00e1nh:" in match_part else "Tranh:"
            if avoid_key in match_part:
                cham_val = match_part.split(avoid_key)[0].strip()
                tranh_val = match_part.split(avoid_key)[1].strip()
                strengths.append(f"Khac che tot: {cham_val}")
                weaknesses.append(f"Can ne tranh: {tranh_val}")
            else:
                strengths.append(f"Khac che tot: {match_part.strip()}")
        if not strengths:
            strengths.append(f"Doi hinh {raw_t['troop']} {raw_t['tier']} chuan meta")
        if not weaknesses:
            weaknesses.append("Can can cu theo binh chung va chien phap doi phuong de linh hoat khac che")
        team_name_str = f"{raw_t['troop']} {team_faction} ({' - '.join(gen_names_key)})"
        team_id = f"meta_team_{counter:03d}_{slugify(team_name_str[:30])}"
        counter += 1
        processed_teams.append({
            "id": team_id,
            "name": team_name_str,
            "tier": raw_t['tier'],
            "season": raw_t.get('season', 'PK'),
            "source": raw_t.get('source', ''),
            "faction": team_faction,
            "troop": raw_t['troop'],
            "description": notes_str if notes_str else f"Doi hinh {team_name_str} xep hang {raw_t['tier']}.",
            "strengths": strengths,
            "weaknesses": weaknesses,
            "generals": team_gens
        })

    tier_order = {"T0": 0, "T0.5": 1, "T1": 2, "T2": 3}
    processed_teams.sort(key=lambda t: (t.get('season', 'PK') != 'PK', tier_order.get(t['tier'], 99)))

    generals_list = sorted(generals_data.values(), key=lambda g: (g['faction'], g['name']))
    tactics_list = sorted(tactics_data.values(), key=lambda t: (t['type'], t['name']))

    print("=" * 50)
    print("BUILD SUMMARY:")
    print(f"  Generals:              {len(generals_list)}")
    print(f"  Tactics:               {len(tactics_list)}")
    print(f"  Total Meta Teams:      {len(processed_teams)}")
    print(f"  Coexisting Portfolios: {len(coexisting_portfolios)}")
    print(f"  Team Do Uy:            {len(team_do_uy_data)}")
    print("=" * 50)

    def save(filename, data):
        path = os.path.join(DB_DIR, filename)
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        size = os.path.getsize(path)
        print(f"  Saved {filename} ({size // 1024}KB)")

    save("generals.json", generals_list)
    save("tactics.json", tactics_list)
    save("meta_teams.json", processed_teams)
    save("coexisting_portfolios.json", coexisting_portfolios)
    save("team_do_uy.json", team_do_uy_data)
    save("cot_truyen.json", cot_truyen_data)
    if isinstance(khai_hoang_data, dict):
        save("khai_hoang.json", [khai_hoang_data])
    else:
        save("khai_hoang.json", khai_hoang_data)
    if isinstance(cai_tao_data, dict):
        save("cai_tao_binh_chung.json", [cai_tao_data])
    else:
        save("cai_tao_binh_chung.json", cai_tao_data)

    print("\nDone!")

if __name__ == "__main__":
    main()
