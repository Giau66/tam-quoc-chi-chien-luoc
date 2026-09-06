# -*- coding: utf-8 -*-
"""
Import and build complete Tam Quoc Chi Chien Luoc database from Excel file:
'Meta Team Giới Thiệu Mùa PK.xlsx'
"""
import openpyxl
import re
import os
import json
import unicodedata

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, "Meta Team Giới Thiệu Mùa PK.xlsx")
DB_DIR = os.path.join(BASE_DIR, "database")

def slugify(text):
    text = unicodedata.normalize('NFD', str(text))
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    text = text.lower().replace('đ', 'd').replace('đ', 'd')
    text = re.sub(r'[^a-z0-9]+', '_', text).strip('_')
    return text

GEN_NORM = {
    'giả hư': 'Giả Hủ',
    'thào tháo': 'Tào Tháo',
    'lăm thống': 'Lăng Thống',
    'sp quãn vũ': 'SP Quan Vũ',
    'spđổng trác': 'SP Đổng Trác',
    'lục tốṇ': 'Lục Tốn',
    'bái sư hứa du': 'Hứa Du',
    'bái sư lỗ túc': 'Lỗ Túc',
    'bái sư pháp chính': 'Pháp Chính',
    'bái sư sp tuân úc': 'SP Tuân Úc',
    'lữ bố 40th': 'Lữ Bố',
    'co tinh thái': 'Trương Tinh Thái',
    'co đại kiều': 'Đại Kiều',
    'co tiểu kiều': 'Tiểu Kiều',
    'tinh thái': 'Trương Tinh Thái',
    'tư mã ý': 'Tư Mã Ý',
    'mãn sủng': 'Mãn Sủng',
    'triệu vân': 'Triệu Vân',
    'khương duy': 'Khương Duy',
    'bàng thống': 'Bàng Thống',
    'gia cát lượng': 'Gia Cát Lượng',
    'tào tháo': 'Tào Tháo',
    'chu thái': 'Chu Thái',
    'lăng thống': 'Lăng Thống',
    'lưu bị': 'Lưu Bị',
    'trương phi': 'Trương Phi',
    'quan vũ': 'Quan Vũ',
    'ngụy diên': 'Ngụy Diên',
    'quan ngân bình': 'Quan Ngân Bình',
    'pháp chính': 'Pháp Chính',
    'sp mã siêu': 'SP Mã Siêu',
    'sp hoàng phổ tung': 'SP Hoàng Phổ Tung',
    'sp tuân úc': 'SP Tuân Úc',
    'sp quách gia': 'SP Quách Gia',
    'sp quan vũ': 'SP Quan Vũ',
    'sp chu du': 'SP Chu Du',
    'sp lữ mông': 'SP Lữ Mông',
    'sp viên thiệu': 'SP Viên Thiệu',
    'sp chu tuấn': 'SP Chu Tuấn',
    'sp đổng trác': 'SP Đổng Trác',
    'sp điêu thuyền': 'SP Điêu Thuyền',
    'sp lô trực': 'SP Lô Trực',
    'sp tôn kiên': 'SP Tôn Kiên',
    'sp hoàng nguyệt anh': 'SP Hoàng Nguyệt Anh',
    'mã quân': 'Mã Quân',
    'sp hứa chử': 'SP Hứa Chử',
    'sp tào chân': 'SP Tào Chân',
    'dương dị': 'Dương Dị',
    'sp trương bảo': 'SP Trương Bảo',
    'sp trương lương': 'SP Trương Lương',
    'sp bàng đức': 'SP Bàng Đức',
    'sp lưu diệp': 'SP Lưu Diệp',
    'gia cát khác': 'Gia Cát Khác',
    'mã đại': 'Mã Đại',
    'sp pháp chính': 'SP Pháp Chính'
}

TAC_NORM = {
    'bát môn kim tọa trận': 'Bát Môn Kim Tỏa Trận',
    'bạch nhi binh': 'Bạch Nhị Binh',
    'cạo xương trị độc': 'Cắt Xương Trị Độc',
    'chờ đời xuất phát': 'Chờ Đợi Xuất Phát',
    'tiềm long trần': 'Tiềm Long Trận',
    'dùng võ thông thần': 'Dụng Võ Thần Thông',
    'hoành tào thiên quân': 'Hoành Tảo Thiên Quân',
    'phong thi trận': 'Phong Thỉ Trận',
    'thuyền cỏ mượn tên': 'Thảo Thuyền Mượn Tên',
    'thác thuyền cỏ mượn tên': 'Thảo Thuyền Mượn Tên',
    'thác·thuyền cỏ mượn tên': 'Thảo Thuyền Mượn Tên',
    'an ủi quân dân': 'Quân Dân Khích Lệ',
    'tạm lánh sóng gió': 'Tạm Thời Tránh Mũi Nhọn',
    'lính giáp mây': 'Đằng Giáp Binh',
    'chuyển phi hùng quân': 'Phi Hùng Quân',
    'chuyển bạch nhị binh': 'Bạch Nhị Binh',
    'tinh lính đan dương': 'Lính Đan Dương',
    'tinh hổ báo kỵ': 'Hổ Báo Kỵ',
    'tinh lính thanh châu': 'Lính Thanh Châu',
    'tinh giải phiền vệ': 'Giải Phiền Vệ',
    'tinh tử sĩ tiên phong': 'Tử Sĩ Tiên Phong',
    'tinh phi hùng quân': 'Phi Hùng Quân',
    'tinh quân hổ vệ': 'Quân Hổ Vệ'
}

FACTION_MAP = {
    # Thục
    'Lưu Bị': 'Thục', 'Quan Vũ': 'Thục', 'Trương Phi': 'Thục', 'Triệu Vân': 'Thục',
    'Gia Cát Lượng': 'Thục', 'Bàng Thống': 'Thục', 'Khương Duy': 'Thục', 'Ngụy Diên': 'Thục',
    'Hoàng Trung': 'Thục', 'Mã Siêu': 'Thục', 'Pháp Chính': 'Thục', 'Quan Ngân Bình': 'Thục',
    'Trương Tinh Thái': 'Thục', 'Hoàng Nguyệt Anh': 'Thục', 'Nghiêm Nhan': 'Thục',
    'Mã Vân Lộc': 'Thục', 'Vương Bình': 'Thục', 'Từ Thứ': 'Thục', 'SP Quan Vũ': 'Thục',
    'SP Gia Cát Lượng': 'Thục', 'SP Hoàng Nguyệt Anh': 'Thục', 'SP Mã Siêu': 'Thục',
    'Mã Đại': 'Thục', 'SP Pháp Chính': 'Thục', 'Liêu Hóa': 'Thục',
    # Ngụy
    'Tào Tháo': 'Ngụy', 'Tư Mã Ý': 'Ngụy', 'Quách Gia': 'Ngụy', 'Giả Hủ': 'Ngụy',
    'Trình Dục': 'Ngụy', 'Hạ Hầu Uyên': 'Ngụy', 'Hạ Hầu Đôn': 'Ngụy', 'Trương Liêu': 'Ngụy',
    'Hác Chiêu': 'Ngụy', 'Mãn Sủng': 'Ngụy', 'Tuân Úc': 'Ngụy', 'Tuân Du': 'Ngụy',
    'Vương Nguyên Cơ': 'Ngụy', 'Chung Hội': 'Ngụy', 'Đặng Ngải': 'Ngụy', 'Hứa Chử': 'Ngụy',
    'Điển Vi': 'Ngụy', 'Tào Nhân': 'Ngụy', 'Tào Phi': 'Ngụy', 'Tào Thuần': 'Ngụy',
    'Bàng Đức': 'Ngụy', 'Chân Cơ': 'Ngụy', 'SP Tuân Úc': 'Ngụy', 'SP Quách Gia': 'Ngụy',
    'SP Hứa Chử': 'Ngụy', 'SP Tào Chân': 'Ngụy', 'SP Bàng Đức': 'Ngụy', 'SP Lưu Diệp': 'Ngụy',
    # Ngô
    'Tôn Quyền': 'Ngô', 'Tôn Sách': 'Ngô', 'Tôn Kiên': 'Ngô', 'Tôn Thượng Hương': 'Ngô',
    'Chu Du': 'Ngô', 'Lục Tốn': 'Ngô', 'Lỗ Túc': 'Ngô', 'Lữ Mông': 'Ngô',
    'Thái Sử Từ': 'Ngô', 'Cam Ninh': 'Ngô', 'Trình Phổ': 'Ngô', 'Hoàng Cái': 'Ngô',
    'Chu Thái': 'Ngô', 'Lăng Thống': 'Ngô', 'Đại Kiều': 'Ngô', 'Tiểu Kiều': 'Ngô',
    'Lục Kháng': 'Ngô', 'Đinh Phụng': 'Ngô', 'Hàn Đương': 'Ngô', 'Tưởng Khâm': 'Ngô',
    'SP Chu Du': 'Ngô', 'SP Lữ Mông': 'Ngô', 'SP Tôn Kiên': 'Ngô', 'Gia Cát Khác': 'Ngô',
    # Quần
    'Lữ Bố': 'Quần', 'Điêu Thuyền': 'Quần', 'Đổng Trác': 'Quần', 'Hoa Đà': 'Quần',
    'Tả Từ': 'Quần', 'Vu Cát': 'Quần', 'Trương Giác': 'Quần', 'Viên Thiệu': 'Quần',
    'Chu Tuấn': 'Quần', 'Thư Thụ': 'Quần', 'Hoa Hùng': 'Quần', 'Nhan Lương': 'Quần',
    'Văn Xú': 'Quần', 'Chúc Dung': 'Quần', 'Mạnh Hoạch': 'Quần', 'Trần Cung': 'Quần',
    'Điền Phong': 'Quần', 'Lý Nho': 'Quần', 'Hứa Du': 'Quần', 'Dương Dị': 'Quần',
    'Mã Quân': 'Quần', 'SP Viên Thiệu': 'Quần', 'SP Chu Tuấn': 'Quần', 'SP Đổng Trác': 'Quần',
    'SP Điêu Thuyền': 'Quần', 'SP Lô Trực': 'Quần', 'SP Hoàng Phổ Tung': 'Quần',
    'SP Trương Bảo': 'Quần', 'SP Trương Lương': 'Quần', 'Khúc Nghĩa': 'Quần',
    'Cao Lãm': 'Quần', 'Chu Thiệu': 'Quần', 'Viên Thuật': 'Quần'
}

def clean_gen_name(name):
    if not name:
        return ""
    name = str(name).strip()
    name = re.sub(r'[\r\n]+', ' ', name)
    name = re.sub(r'\([^\)]*\)', '', name)
    name = re.sub(r'\b(40th|CO|Chủ Tướng|Phó Tướng)\b', '', name, flags=re.I)
    name = name.replace('̣', '').strip()
    name = re.sub(r'\s+', ' ', name)
    n_lower = name.lower()
    return GEN_NORM.get(n_lower, name)

def clean_tac_name(name):
    if not name:
        return ""
    name = str(name).strip()
    name = re.sub(r'[\r\n]+', ' ', name)
    name = re.sub(r'^(Tinh|Chuyển|Thác)[·\.\s]', '', name)
    name = re.sub(r'\([^\)]*\)', '', name)
    name = name.strip()
    name = re.sub(r'\s+', ' ', name)
    n_lower = name.lower()
    return TAC_NORM.get(n_lower, name)

def infer_faction(name):
    clean = clean_gen_name(name)
    if clean in FACTION_MAP:
        return FACTION_MAP[clean]
    for k, v in FACTION_MAP.items():
        if clean == k or clean in k:
            return v
    return "Quần"

def infer_tactic_type(name):
    n = name.lower()
    if 'trận' in n:
        return 'Trận pháp'
    if any(k in n for k in ['kỵ', 'khiên', 'cung', 'thương', 'binh', 'doanh', 'vệ', 'sĩ']):
        return 'Binh chủng'
    if any(k in n for k in ['thái bình', 'sĩ biệt', 'dụng võ', 'tuyệt địa', 'hổ cứ', 'lõa y', 'binh vô', 'văn thao']):
        return 'Bị động'
    if any(k in n for k in ['bát môn', 'quân dân', 'thảo thuyền', 'tạm thời', 'thịnh khí', 'tiềm long', 'cơ hình', 'giải phiền']):
        return 'Chỉ huy'
    if any(k in n for k in ['nhất kỵ', 'quỷ thần', 'bách kỵ', 'bạo lệ', 'đương phong', 'mau giành']):
        return 'Đột kích'
    return 'Chủ động'

def main():
    print(f"Loading Excel file from {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    generals_data = {}
    tactics_data = {}
    meta_teams_list = []
    coexisting_portfolios = []

    # -------------------------------------------------------------
    # 1. PARSE Tướng&CP SHEET
    # -------------------------------------------------------------
    if 'Tướng&CP' in wb.sheetnames:
        s_cp = wb['Tướng&CP']
        mode = None
        season = "PK"
        for r in range(1, s_cp.max_row + 1):
            c1 = str(s_cp.cell(r, 1).value or '').strip()
            c2 = str(s_cp.cell(r, 2).value or '').strip()
            c3 = str(s_cp.cell(r, 3).value or '').strip()
            c4 = str(s_cp.cell(r, 4).value or '').strip()
            c5 = str(s_cp.cell(r, 5).value or '').strip()
            
            if 'PK' in c1 or 'mùa' in c1.lower():
                season = c1
            if 'Tướng Mới' in c1 or 'Tướng Mới' in c2:
                mode = 'GEN'
                continue
            elif 'Chiến Pháp Sự Kiện' in c1 or 'Chiến Pháp Sự Kiện' in c2 or 'Chiến Pháp Kế Thừa' in c2:
                mode = 'TAC'
                continue
            if c1 in ['STT', ''] and c2 in ['Tên Tướng', 'Tên Chiến Pháp']:
                continue
                
            if mode == 'GEN' and c2:
                g_name = clean_gen_name(c2)
                if g_name:
                    f = infer_faction(g_name)
                    gid = f"{slugify(f)}_{slugify(g_name)}"
                    generals_data[g_name] = {
                        "id": gid,
                        "name": g_name,
                        "faction": f,
                        "cost": 7 if 'SP' in g_name else 6,
                        "troop": {"Kỵ": "S", "Khiên": "A", "Cung": "S", "Thương": "S", "Khí": "A"},
                        "inherent_skill": c3 if c3 else f"Thiên phú của {g_name}",
                        "inherit_tactic": clean_tac_name(c4) if c4 != '-' else '',
                        "role": [c5] if c5 and c5 != '-' else ["Chủ lực / Hỗ trợ chiến thuật"],
                        "season": season
                    }
                if c4 and c4 != '-':
                    t_name = clean_tac_name(c4)
                    if t_name:
                        tid = slugify(t_name)
                        tactics_data[t_name] = {
                            "id": tid,
                            "name": t_name,
                            "type": infer_tactic_type(t_name),
                            "quality": "S",
                            "description": f"Chiến pháp kế thừa từ {g_name}",
                            "season": season
                        }
            elif mode == 'TAC' and c2:
                t_name = clean_tac_name(c2)
                if t_name:
                    tid = slugify(t_name)
                    desc = c3
                    if c5:
                        desc += f" (Cần: {c5})"
                    tactics_data[t_name] = {
                        "id": tid,
                        "name": t_name,
                        "type": infer_tactic_type(t_name),
                        "quality": "S",
                        "description": desc,
                        "suitable_generals": c4 if c4 else "",
                        "season": season
                    }

    # -------------------------------------------------------------
    # 2. PARSE Team giới thiệu(mới)
    # -------------------------------------------------------------
    if 'Team giới thiệu(mới)' in wb.sheetnames:
        s_new = wb['Team giới thiệu(mới)']
        cur_t = None
        for r in range(5, s_new.max_row + 1):
            type_val = str(s_new.cell(r, 1).value or '').strip()
            gen_val = str(s_new.cell(r, 2).value or '').strip()
            cp1_val = str(s_new.cell(r, 3).value or '').strip()
            cp2_val = str(s_new.cell(r, 4).value or '').strip()
            bt1 = str(s_new.cell(r, 5).value or '').strip()
            bt2 = str(s_new.cell(r, 6).value or '').strip()
            bt3 = str(s_new.cell(r, 7).value or '').strip()
            note = str(s_new.cell(r, 8).value or '').strip()
            
            if type_val:
                if cur_t and len(cur_t['generals']) >= 2:
                    meta_teams_list.append(cur_t)
                troop = "Kỵ"
                for t in ["Thương", "Khiên", "Cung", "Kỵ", "Khí"]:
                    if t in type_val:
                        troop = t
                        break
                tier = "T0"
                if "⭐⭐⭐⭐" in type_val or "★★★★" in type_val:
                    tier = "T0"
                elif "⭐⭐⭐" in type_val or "★★★" in type_val:
                    tier = "T0.5"
                elif "⭐⭐" in type_val or "★★" in type_val:
                    tier = "T1"
                cur_t = {
                    "source": "Team giới thiệu(mới)",
                    "troop": troop,
                    "tier": tier,
                    "note": note,
                    "generals": []
                }
            if cur_t and gen_val:
                cur_t['generals'].append({
                    "raw_name": gen_val,
                    "cp1_raw": cp1_val,
                    "cp2_raw": cp2_val,
                    "binh_thu": [b for b in [bt1, bt2, bt3] if b]
                })
                if note and not cur_t['note']:
                    cur_t['note'] = note
        if cur_t and len(cur_t['generals']) >= 2:
            meta_teams_list.append(cur_t)

    # -------------------------------------------------------------
    # 3. PARSE Meta Team (Official meta rankings)
    # -------------------------------------------------------------
    if 'Meta Team' in wb.sheetnames:
        s_meta = wb['Meta Team']
        cur_t = None
        for r in range(3, s_meta.max_row + 1):
            stt = str(s_meta.cell(r, 1).value or '').strip()
            type_val = str(s_meta.cell(r, 2).value or '').strip()
            power = str(s_meta.cell(r, 3).value or '').strip()
            gen_val = str(s_meta.cell(r, 4).value or '').strip()
            cp1_val = str(s_meta.cell(r, 5).value or '').strip()
            cp2_val = str(s_meta.cell(r, 6).value or '').strip()
            note = str(s_meta.cell(r, 10).value or '').strip()
            
            if stt or type_val:
                if cur_t and len(cur_t['generals']) >= 2:
                    meta_teams_list.append(cur_t)
                troop = "Kỵ"
                for t in ["Thương", "Khiên", "Cung", "Kỵ", "Khí"]:
                    if t in type_val:
                        troop = t
                        break
                cur_t = {
                    "source": "Meta Team",
                    "troop": troop,
                    "tier": power if power else "T1",
                    "note": note,
                    "generals": []
                }
            if cur_t and gen_val:
                cur_t['generals'].append({
                    "raw_name": gen_val,
                    "cp1_raw": cp1_val,
                    "cp2_raw": cp2_val,
                    "binh_thu": []
                })
                if note and not cur_t['note']:
                    cur_t['note'] = note
        if cur_t and len(cur_t['generals']) >= 2:
            meta_teams_list.append(cur_t)

    # -------------------------------------------------------------
    # 4. PARSE Meta mỗi phe (cũ)
    # -------------------------------------------------------------
    if 'Meta mỗi phe (cũ)' in wb.sheetnames:
        s_phe = wb['Meta mỗi phe (cũ)']
        cur_t = None
        for r in range(4, s_phe.max_row + 1):
            type_val = str(s_phe.cell(r, 1).value or '').strip()
            gen_val = str(s_phe.cell(r, 2).value or '').strip()
            cp1_val = str(s_phe.cell(r, 3).value or '').strip()
            cp2_val = str(s_phe.cell(r, 4).value or '').strip()
            note = str(s_phe.cell(r, 8).value or '').strip()
            
            if type_val:
                if cur_t and len(cur_t['generals']) >= 2:
                    meta_teams_list.append(cur_t)
                troop = "Kỵ"
                for t in ["Thương", "Khiên", "Cung", "Kỵ", "Khí"]:
                    if t in type_val:
                        troop = t
                        break
                cur_t = {
                    "source": "Meta mỗi phe (cũ)",
                    "troop": troop,
                    "tier": "T1",
                    "note": note,
                    "generals": []
                }
            if cur_t and gen_val:
                cur_t['generals'].append({
                    "raw_name": gen_val,
                    "cp1_raw": cp1_val,
                    "cp2_raw": cp2_val,
                    "binh_thu": []
                })
                if note and not cur_t['note']:
                    cur_t['note'] = note
        if cur_t and len(cur_t['generals']) >= 2:
            meta_teams_list.append(cur_t)

    # -------------------------------------------------------------
    # 5. PARSE Coexisting Portfolios (Team Cùng Tồn & nhiều Team cùng tồn tại)
    # -------------------------------------------------------------
    if 'Team Cùng Tồn' in wb.sheetnames:
        s_ct = wb['Team Cùng Tồn']
        cur_set = None
        cur_sub_team = None
        for r in range(4, s_ct.max_row + 1):
            c1 = str(s_ct.cell(r, 1).value or '').strip()
            c2 = str(s_ct.cell(r, 2).value or '').strip()
            c3 = str(s_ct.cell(r, 3).value or '').strip()
            c4 = str(s_ct.cell(r, 4).value or '').strip()
            c5 = str(s_ct.cell(r, 5).value or '').strip()
            
            if 'Cùng tồn' in c1:
                if cur_set:
                    coexisting_portfolios.append(cur_set)
                cur_set = {
                    "set_name": c1,
                    "description": c2,
                    "teams": []
                }
                cur_sub_team = None
                continue
                
            if c2 and ('19C' in c2 or '20C' in c2 or '18C' in c2 or '16C' in c2 or 'Thương' in c2 or 'Khiên' in c2 or 'Cung' in c2 or 'Kỵ' in c2):
                cur_sub_team = {
                    "team_name": c2,
                    "generals": []
                }
                if cur_set:
                    cur_set["teams"].append(cur_sub_team)
                    
            if cur_sub_team and c3:
                g_c = clean_gen_name(c3)
                if g_c and g_c not in ['Võ Tướng', 'Tướng']:
                    cur_sub_team["generals"].append({
                        "name": g_c,
                        "tactic1": clean_tac_name(c4),
                        "tactic2": clean_tac_name(c5)
                    })
        if cur_set:
            coexisting_portfolios.append(cur_set)

    # -------------------------------------------------------------
    # 6. DEDUPLICATE & STRUCTURE META TEAMS
    # -------------------------------------------------------------
    processed_teams = []
    seen_team_keys = set()
    team_counter = 1

    positions = ["Chủ tướng", "Phó tướng 1", "Phó tướng 2"]

    for raw_t in meta_teams_list:
        gens = raw_t.get('generals', [])
        if len(gens) < 2:
            continue
            
        team_gens = []
        gen_names_for_key = []
        
        for idx, g_obj in enumerate(gens[:3]):
            raw_name = g_obj['raw_name']
            # Alternatives separated by '/'
            alt_candidates = [clean_gen_name(p) for p in raw_name.split('/') if clean_gen_name(p)]
            if not alt_candidates:
                continue
            main_gen_name = alt_candidates[0]
            alt_gens = alt_candidates[1:]
            
            gen_names_for_key.append(main_gen_name)
            
            # Register general in generals_data if missing
            if main_gen_name not in generals_data:
                f = infer_faction(main_gen_name)
                gid = f"{slugify(f)}_{slugify(main_gen_name)}"
                generals_data[main_gen_name] = {
                    "id": gid,
                    "name": main_gen_name,
                    "faction": f,
                    "cost": 7 if 'SP' in main_gen_name else 6,
                    "troop": {"Kỵ": "A", "Khiên": "A", "Cung": "A", "Thương": "A", "Khí": "B"},
                    "inherent_skill": f"Kỹ năng riêng của {main_gen_name}",
                    "role": ["Tướng chủ lực / Hỗ trợ trong Meta"],
                    "season": "PK"
                }
                
            # Parse tactics
            bis_tactics = []
            sub_tactics = []
            for cp_raw in [g_obj.get('cp1_raw', ''), g_obj.get('cp2_raw', '')]:
                if not cp_raw:
                    continue
                parts = [clean_tac_name(p) for p in re.split(r'[/,\n]+', cp_raw) if clean_tac_name(p)]
                if parts:
                    bis_tactics.append(parts[0])
                    # Register tactic if missing
                    if parts[0] not in tactics_data:
                        tactics_data[parts[0]] = {
                            "id": slugify(parts[0]),
                            "name": parts[0],
                            "type": infer_tactic_type(parts[0]),
                            "quality": "S",
                            "description": f"Chiến pháp mạnh được khuyên dùng cho {main_gen_name}",
                            "season": "PK"
                        }
                    for alt_t in parts[1:]:
                        sub_tactics.append(alt_t)
                        if alt_t not in tactics_data:
                            tactics_data[alt_t] = {
                                "id": slugify(alt_t),
                                "name": alt_t,
                                "type": infer_tactic_type(alt_t),
                                "quality": "S",
                                "description": f"Chiến pháp thay thế hiệu quả cho {main_gen_name}",
                                "season": "PK"
                            }
                            
            f = infer_faction(main_gen_name)
            gid = generals_data[main_gen_name]["id"]
            
            team_gens.append({
                "position": positions[idx] if idx < len(positions) else f"Phó tướng {idx}",
                "general_id": gid,
                "name": main_gen_name,
                "bis_tactics": bis_tactics,
                "sub_tactics": sub_tactics,
                "binh_thu": g_obj.get('binh_thu', []),
                "alt_generals": alt_gens
            })
            
        if len(team_gens) < 2:
            continue
            
        # Deduplication key based on sorted general names + troop
        team_key = "_".join(sorted(gen_names_for_key)) + f"_{raw_t['troop']}"
        if team_key in seen_team_keys:
            continue
        seen_team_keys.add(team_key)
        
        # Determine team faction:
        factions = [infer_faction(g['name']) for g in team_gens]
        if len(set(factions)) == 1:
            team_faction = factions[0]
        else:
            team_faction = "Tam Thế"
            
        # Parse notes for strengths / counters
        notes_str = raw_t.get('note', '')
        strengths = []
        weaknesses = []
        
        if "Chạm:" in notes_str:
            match_part = notes_str.split("Chạm:")[1]
            if "Tránh:" in match_part:
                cham_val = match_part.split("Tránh:")[0].strip()
                tranh_val = match_part.split("Tránh:")[1].strip()
                strengths.append(f"Khắc chế tốt: {cham_val}")
                weaknesses.append(f"Cần né tránh: {tranh_val}")
            else:
                strengths.append(f"Khắc chế tốt: {match_part.strip()}")
        elif "Khắc chế" in notes_str:
            strengths.append(notes_str[:120])
            
        if not strengths:
            strengths.append(f"Đội hình {raw_t['troop']} {raw_t['tier']} chuẩn meta với sức mạnh tối ưu")
        if not weaknesses:
            weaknesses.append("Cần căn cứ theo binh chủng và chiến pháp đối phương để linh hoạt khắc chế")
            
        team_name_str = f"{raw_t['troop']} {team_faction} ({' - '.join(gen_names_for_key)})"
        
        team_id = f"meta_team_{team_counter:03d}_{slugify(team_name_str[:30])}"
        team_counter += 1
        
        processed_teams.append({
            "id": team_id,
            "name": team_name_str,
            "tier": raw_t['tier'],
            "season": "PK",
            "faction": team_faction,
            "troop": raw_t['troop'],
            "description": notes_str if notes_str else f"Đội hình {team_name_str} xếp hạng {raw_t['tier']} với các chiến pháp phối hợp hiệu quả cao.",
            "strengths": strengths,
            "weaknesses": weaknesses,
            "generals": team_gens
        })

    # Sort meta teams by Tier (T0 -> T0.5 -> T1)
    tier_order = {"T0": 0, "T0.5": 1, "T1": 2, "T2": 3}
    processed_teams.sort(key=lambda t: tier_order.get(t['tier'], 99))

    # Convert generals_data and tactics_data to lists
    generals_list = list(generals_data.values())
    generals_list.sort(key=lambda g: (g['faction'], g['name']))
    
    tactics_list = list(tactics_data.values())
    tactics_list.sort(key=lambda t: (t['type'], t['name']))

    print(f"==================================================")
    print(f"BUILD SUMMARY FROM EXCEL:")
    print(f"  * Total Generals: {len(generals_list)}")
    print(f"  * Total Tactics:  {len(tactics_list)}")
    print(f"  * Total Meta Teams: {len(processed_teams)}")
    print(f"  * Coexisting Portfolios: {len(coexisting_portfolios)}")
    print(f"==================================================")

    # Write to database files
    with open(os.path.join(DB_DIR, "generals.json"), "w", encoding="utf-8") as f:
        json.dump(generals_list, f, ensure_ascii=False, indent=2)
    print(" Saved database/generals.json")

    with open(os.path.join(DB_DIR, "tactics.json"), "w", encoding="utf-8") as f:
        json.dump(tactics_list, f, ensure_ascii=False, indent=2)
    print(" Saved database/tactics.json")

    with open(os.path.join(DB_DIR, "meta_teams.json"), "w", encoding="utf-8") as f:
        json.dump(processed_teams, f, ensure_ascii=False, indent=2)
    print(" Saved database/meta_teams.json")

    with open(os.path.join(DB_DIR, "coexisting_portfolios.json"), "w", encoding="utf-8") as f:
        json.dump(coexisting_portfolios, f, ensure_ascii=False, indent=2)
    print(" Saved database/coexisting_portfolios.json")

if __name__ == "__main__":
    main()
