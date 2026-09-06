# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

"""
Script bổ sung: Parse các sheet mới từ Excel và tạo seasons.json
Chạy sau khi import_from_excel.py đã được chạy trước đó.
Sheets mới: Meta team mùa 4, Meta team mùa 5, Team Đô Úy, seasons metadata
"""
import openpyxl
import json
import os
import re
import unicodedata

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, "Meta Team Giới Thiệu Mùa PK.xlsx")
DB_DIR = os.path.join(BASE_DIR, "database")
HINH_DIR = os.path.join(BASE_DIR, "hinh")

def slugify(text):
    text = unicodedata.normalize('NFD', str(text))
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    text = text.lower().replace('đ', 'd')
    text = re.sub(r'[^a-z0-9]+', '_', text).strip('_')
    return text

def clean_cell(v):
    if v is None:
        return ''
    s = str(v).strip()
    s = re.sub(r'[\r\n]+', ' ', s).strip()
    return s

def parse_season_sheet(ws, start_row, season_label, tier_default='T1'):
    """Generic parser for season team sheets (Mùa 4, Mùa 5, Đô Úy)."""
    teams = []
    cur_t = None
    for r in range(start_row, ws.max_row + 1):
        type_val = clean_cell(ws.cell(r, 1).value)
        gen_val  = clean_cell(ws.cell(r, 2).value)
        cp1_val  = clean_cell(ws.cell(r, 3).value)
        cp2_val  = clean_cell(ws.cell(r, 4).value)
        bt1      = clean_cell(ws.cell(r, 5).value)
        bt2      = clean_cell(ws.cell(r, 6).value)
        bt3      = clean_cell(ws.cell(r, 7).value)
        note     = clean_cell(ws.cell(r, 8).value)

        if type_val and type_val not in ('Type', 'STT', 'Thêm'):
            if cur_t and cur_t['generals']:
                teams.append(cur_t)
            troop = 'Kỵ'
            for t in ['Thương', 'Khiên', 'Cung', 'Kỵ', 'Khí']:
                if t in type_val:
                    troop = t
                    break
            cur_t = {
                'season': season_label,
                'troop': troop,
                'tier': tier_default,
                'note': note,
                'generals': []
            }
        if cur_t and gen_val and gen_val not in ('Tướng', 'Võ Tướng', ''):
            cur_t['generals'].append({
                'name': gen_val,
                'cp1': cp1_val,
                'cp2': cp2_val,
                'binh_thu': [b for b in [bt1, bt2, bt3] if b and b != '-']
            })
            if note and not cur_t['note']:
                cur_t['note'] = note

    if cur_t and cur_t['generals']:
        teams.append(cur_t)
    return teams

def main():
    print(f"Loading Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    # -------------------------------------------------------
    # 1. Parse season team sheets
    # -------------------------------------------------------
    season_teams = {}

    # --- PK20 teams are already in meta_teams.json as "PK" season
    # We tag them as pk20 for the season page
    existing_meta = []
    with open(os.path.join(DB_DIR, "meta_teams.json"), "r", encoding="utf-8") as f:
        existing_meta = json.load(f)

    pk20_teams = []
    for t in existing_meta:
        t_copy = dict(t)
        t_copy['season_id'] = 'pk20'
        pk20_teams.append(t_copy)
    season_teams['pk20'] = pk20_teams

    # --- Mùa 4
    mua4_teams = []
    if 'Meta team mùa 4' in wb.sheetnames:
        ws4 = wb['Meta team mùa 4']
        raw = parse_season_sheet(ws4, start_row=3, season_label='Mùa 4', tier_default='T0')
        for idx, t in enumerate(raw):
            if len(t['generals']) < 2:
                continue
            g_names = [g['name'] for g in t['generals']]
            team_id = f"mua4_{idx+1:03d}_{slugify('_'.join(g_names[:2]))}"
            mua4_teams.append({
                'id': team_id,
                'name': f"{t['troop']} ({' - '.join(g_names)})",
                'season_id': 'mua4',
                'troop': t['troop'],
                'tier': t['tier'],
                'note': t['note'],
                'generals': t['generals']
            })
    season_teams['mua4'] = mua4_teams

    # --- Mùa 5
    mua5_teams = []
    if 'Meta team mùa 5' in wb.sheetnames:
        ws5 = wb['Meta team mùa 5']
        raw = parse_season_sheet(ws5, start_row=2, season_label='Mùa 5', tier_default='T0')
        for idx, t in enumerate(raw):
            if len(t['generals']) < 2:
                continue
            g_names = [g['name'] for g in t['generals']]
            team_id = f"mua5_{idx+1:03d}_{slugify('_'.join(g_names[:2]))}"
            mua5_teams.append({
                'id': team_id,
                'name': f"{t['troop']} ({' - '.join(g_names)})",
                'season_id': 'mua5',
                'troop': t['troop'],
                'tier': t['tier'],
                'note': t['note'],
                'generals': t['generals']
            })
    season_teams['mua5'] = mua5_teams

    # --- Team Đô Úy (PK20 mode-specific)
    douy_teams = []
    if 'Team Đô Úy' in wb.sheetnames:
        wsd = wb['Team Đô Úy']
        raw = parse_season_sheet(wsd, start_row=6, season_label='PK20', tier_default='T0.5')
        for idx, t in enumerate(raw):
            if not t['generals']:
                continue
            g_names = [g['name'] for g in t['generals']]
            douy_teams.append({
                'id': f"douy_{idx+1:03d}",
                'name': f"{t['troop']} ({' - '.join(g_names)})",
                'troop': t['troop'],
                'tier': t['tier'],
                'note': t['note'],
                'generals': t['generals']
            })

    # --- Team Giới Thiệu mới (PK20)
    intro_teams = []
    if 'Team giới thiệu(mới)' in wb.sheetnames:
        wsi = wb['Team giới thiệu(mới)']
        raw = parse_season_sheet(wsi, start_row=5, season_label='PK20', tier_default='T0')
        for idx, t in enumerate(raw):
            if not t['generals']:
                continue
            g_names = [g['name'] for g in t['generals']]
            intro_teams.append({
                'id': f"intro_{idx+1:03d}",
                'name': f"{t['troop']} ({' - '.join(g_names)})",
                'troop': t['troop'],
                'tier': t['tier'],
                'note': t['note'],
                'generals': t['generals']
            })

    # New generals from PK20 Tướng&CP
    pk20_new_gens = []
    if 'Tướng&CP' in wb.sheetnames:
        scp = wb['Tướng&CP']
        mode = None
        for r in range(1, scp.max_row + 1):
            c1 = clean_cell(scp.cell(r, 1).value)
            c2 = clean_cell(scp.cell(r, 2).value)
            c3 = clean_cell(scp.cell(r, 3).value)
            c4 = clean_cell(scp.cell(r, 4).value)
            c5 = clean_cell(scp.cell(r, 5).value)
            if 'Tướng Mới' in c1 or 'Tướng Mới' in c2:
                mode = 'GEN'
                continue
            if c1 in ('STT', '') and c2 in ('Tên Tướng', ''):
                continue
            if mode == 'GEN' and c2 and c2 not in ('', 'STT', 'Tên Tướng'):
                pk20_new_gens.append({
                    'name': c2,
                    'skill': c3,
                    'inherit_tactic': c4 if c4 != '-' else '',
                    'role': c5 if c5 != '-' else ''
                })

    # -------------------------------------------------------
    # 2. Get all images from hinh/ folder
    # -------------------------------------------------------
    all_images = []
    if os.path.isdir(HINH_DIR):
        all_images = sorted([
            f for f in os.listdir(HINH_DIR)
            if f.lower().endswith(('.jpg', '.jpeg', '.png', '.webp'))
        ])
    print(f"Found {len(all_images)} images in hinh/")

    # -------------------------------------------------------
    # 3. Build seasons.json
    # -------------------------------------------------------
    seasons = [
        {
            "id": "pk20",
            "name": "PK20 - Khói Lửa Loạn Thế",
            "short_name": "PK20",
            "description": "Mùa giải PK20 với các tướng mới: SP Tôn Kiên, SP Hoàng Nguyệt Anh, Mã Quân. Xuất hiện đội hình Ngô Cung mạnh mẽ và các phương án phối hợp đa dạng.",
            "is_current": True,
            "images": all_images,
            "new_generals": pk20_new_gens,
            "intro_teams": intro_teams,
            "douy_teams": douy_teams,
            "meta_team_ids": [t['id'] for t in pk20_teams[:30]]  # top 30 meta teams
        },
        {
            "id": "mua5",
            "name": "Mùa 5 - Tam Phân Thiên Hạ",
            "short_name": "Mùa 5",
            "description": "Mùa giải 5 với đội hình Cung SP Viên Thiệu mạnh và nhiều đội hình Khiên phòng thủ.",
            "is_current": False,
            "images": [],
            "new_generals": [],
            "intro_teams": [],
            "douy_teams": [],
            "meta_teams": mua5_teams
        },
        {
            "id": "mua4",
            "name": "Mùa 4 - Hùng Bá Nhất Phương",
            "short_name": "Mùa 4",
            "description": "Mùa giải 4 với đội hình Cung GCL kinh điển (Khương Duy - Gia Cát Lượng - Bàng Thống).",
            "is_current": False,
            "images": [],
            "new_generals": [],
            "intro_teams": [],
            "douy_teams": [],
            "meta_teams": mua4_teams
        }
    ]

    out_path = os.path.join(DB_DIR, "seasons.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(seasons, f, ensure_ascii=False, indent=2)
    print(f"Saved seasons.json ({len(seasons)} seasons, {len(all_images)} images)")
    print(f"  PK20: {len(intro_teams)} intro teams, {len(douy_teams)} đô úy teams")
    print(f"  Mùa 5: {len(mua5_teams)} teams")
    print(f"  Mùa 4: {len(mua4_teams)} teams")

if __name__ == "__main__":
    main()
